""".doc → .docx 转换:LibreOffice 命令形状、临时目录 / profile 生命周期、各类失败。

不依赖本机装没装 LibreOffice:子进程边界(subprocess.run)由测试替身接管,
命令仍是真实拼的,临时目录 / 串行锁 / 输出读回都是真执行。真机集成用例在末尾,
只在装了 soffice 的机器(开发者本机 / ECS)上跑,没装的机器直接跳过。
"""

from __future__ import annotations

import io
import subprocess
import tempfile
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.rag import doc_convert
from app.rag.extract import extract
from tests import sofficekit


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class FakeRun:
    """subprocess.run 的替身:记录命令,按状态产出输出文件 / 返回码 / 抛错。"""

    def __init__(self) -> None:
        self.calls: list[list[str]] = []
        self.output: bytes | None = _docx_bytes("转换后的正文")
        self.rc = 0
        self.stderr = b""
        self.raise_exc: BaseException | None = None
        self.lock_held_at_run: bool | None = None

    def __call__(self, cmd, capture_output=True, timeout=None):  # noqa: ARG002 签名对齐真 subprocess.run
        self.calls.append(list(cmd))
        self.lock_held_at_run = doc_convert._convert_lock.locked()
        if self.raise_exc is not None:
            raise self.raise_exc
        if self.output is not None:
            outdir = Path(cmd[cmd.index("--outdir") + 1])
            (outdir / "src.docx").write_bytes(self.output)
        return SimpleNamespace(returncode=self.rc, stdout=b"", stderr=self.stderr)


@pytest.fixture
def fake(_no_real_libreoffice, monkeypatch):
    # 显式依赖离线夹具,保证这次 setattr 排在它之后(否则会被它覆盖成"未安装")。
    f = FakeRun()
    monkeypatch.setattr(doc_convert, "soffice_bin", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(doc_convert.subprocess, "run", f)
    return f


def _temp_leftovers() -> set[str]:
    return {p.name for p in Path(tempfile.gettempdir()).glob("qa_doc_*")}


# ---- 未安装 / 成功路径 ----

def test_missing_binary_marks_unsupported(monkeypatch):
    monkeypatch.setattr(doc_convert, "soffice_bin", lambda: "")
    docx, err = doc_convert.doc_to_docx(b"old doc bytes")
    assert docx is None
    assert err.startswith("unsupported")
    assert "LibreOffice" in err and "另存为 .docx" in err


def test_command_shape_and_serialized_by_lock(fake):
    before = _temp_leftovers()
    docx, err = doc_convert.doc_to_docx(b"old doc bytes")

    assert err is None and docx == fake.output
    (cmd,) = fake.calls
    assert cmd[0] == "/usr/bin/soffice"
    assert "--headless" in cmd and "--norestore" in cmd
    assert cmd[cmd.index("--convert-to") + 1] == "docx"
    outdir = cmd[cmd.index("--outdir") + 1]
    assert outdir.startswith(tempfile.gettempdir())
    env_arg = next(a for a in cmd if a.startswith("-env:UserInstallation="))
    assert env_arg.removeprefix("-env:UserInstallation=").startswith(Path(outdir).as_uri())
    assert cmd[-1] == str(Path(outdir) / "src.doc")   # 固定 ASCII 名,原地写进临时目录
    assert fake.lock_held_at_run is True              # 转换期间持锁(串行化)
    assert _temp_leftovers() == before                # 临时目录整棵清掉


def test_fresh_profile_per_call(fake):
    doc_convert.doc_to_docx(b"first")
    doc_convert.doc_to_docx(b"second")
    profiles = [a for c in fake.calls for a in c if a.startswith("-env:UserInstallation=")]
    assert len(profiles) == 2 and profiles[0] != profiles[1]


# ---- 失败路径(都返回原因字符串,不抛) ----

def test_nonzero_exit_reports_stderr(fake):
    fake.rc = 1
    fake.stderr = b"Error: source file could not be loaded\n"
    docx, err = doc_convert.doc_to_docx(b"broken")
    assert docx is None
    assert "转换失败" in err and "退出码 1" in err and "source file could not be loaded" in err


def test_missing_output_file_is_error(fake):
    fake.output = None                     # 退出码 0 但没写出 docx(如加密文档被拒)
    docx, err = doc_convert.doc_to_docx(b"no output")
    assert docx is None
    assert "转换失败" in err and "退出码 0" in err   # 报 LibreOffice,不是读文件时的 FileNotFoundError


def test_empty_output_is_error(fake):
    fake.output = b""
    docx, err = doc_convert.doc_to_docx(b"empty output")
    assert docx is None and "为空" in err


def test_timeout_reports_error_and_cleans_up(fake):
    before = _temp_leftovers()
    fake.raise_exc = subprocess.TimeoutExpired("soffice", doc_convert.CONVERT_TIMEOUT_SECONDS)
    docx, err = doc_convert.doc_to_docx(b"slow")
    assert docx is None and "超时" in err
    assert _temp_leftovers() == before     # 超时被杀也要清掉临时目录


def test_launch_failure_reports_error(fake):
    fake.raise_exc = OSError("No such file or directory")
    docx, err = doc_convert.doc_to_docx(b"anything")
    assert docx is None
    assert "无法启动" in err and "No such file" in err


# ---- 经 extract / extract_document 的分派 ----

def test_extract_routes_doc_through_conversion(monkeypatch):
    seen: list[bytes] = []

    def fake_convert(data: bytes):
        seen.append(data)
        return _docx_bytes("转换后的正文"), None

    monkeypatch.setattr(doc_convert, "doc_to_docx", fake_convert)
    res = extract(b"raw-doc-bytes", "老文件.doc")
    assert seen == [b"raw-doc-bytes"]
    assert "转换后的正文" in res.text
    assert res.extractor == "libreoffice+python-docx" and res.error is None


def test_extract_reports_conversion_error(monkeypatch):
    monkeypatch.setattr(
        doc_convert, "doc_to_docx",
        lambda data: (None, "unsupported: .doc 老格式需 LibreOffice 转换,当前机器未安装"),
    )
    res = extract(b"raw", "老文件.doc")
    assert res.text == "" and res.extractor == "none"
    assert res.error.startswith("unsupported")


def test_document_extract_keeps_native_path_for_doc(monkeypatch):
    """.doc 转换后仍走 Office 原生路径:选 OCR 方式时给同样的「不支持 OCR」告警。"""
    from app.rag.document_extract import extract_document

    monkeypatch.setattr(doc_convert, "doc_to_docx", lambda data: (_docx_bytes("会议纪要"), None))
    res = extract_document(b"raw", "会议纪要.doc", "general")
    assert "会议纪要" in res.text
    assert res.method == "native"
    assert any("不支持 OCR" in w for w in res.warnings)


# ---- 真机集成(本机没装 soffice 就跳过) ----

@pytest.mark.skipif(not sofficekit.REAL_SOFFICE, reason="本机没有 LibreOffice(soffice)")
def test_real_libreoffice_roundtrip(tmp_path):
    """docx → .doc(真 soffice) → extract 抽回同样文字;验证的是真命令、真转换。"""
    from docx import Document

    d = Document()
    d.add_paragraph("真实转换往返:报销单编号 A-1")
    seed = tmp_path / "seed.docx"
    d.save(str(seed))
    doc = sofficekit.convert(seed, "doc", tmp_path)
    res = extract(doc.read_bytes(), "seed.doc")
    assert res.error is None, res.error
    assert "报销单编号 A-1" in res.text
