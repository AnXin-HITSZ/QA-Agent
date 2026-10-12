"""Document structure, chunk coverage and published evidence: fully offline."""
from __future__ import annotations

import io
from types import SimpleNamespace

import pytest

from app.config import Settings, get_settings
from app.rag import ingest, retrieve, store
from app.rag.chunking import pack_units, token_bound
from app.rag.document_extract import TextUnit, extract_document
from app.rag.extract import extract
from tests.pdfkit import make_pdf


def _bytes(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_word_keeps_interleaved_table_and_heading_context():
    from docx import Document
    doc = Document()
    doc.add_heading("设备外借", 1)
    doc.add_paragraph("表格之前")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "设备", "管理员"
    table.cell(1, 0).text, table.cell(1, 1).text = "示波器", "张三"
    doc.add_paragraph("表格之后")
    result = extract_document(_bytes(doc), "设备.docx", "native_only")
    text = result.text
    assert text.index("表格之前") < text.index("示波器") < text.index("表格之后")
    rows = [u for u in result.units if u.kind == "table"]
    assert "设备: 示波器" in rows[-1].text
    assert rows[-1].section[0] == "设备外借"


def test_xlsx_merged_title_and_multilevel_headers():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "预算"
    ws.append(["年度预算", None, None])
    ws.merge_cells("A1:C1")
    ws.append(["项目", "经费", None])
    ws.merge_cells("A2:A3")
    ws.merge_cells("B2:C2")
    ws.append([None, "预算", "支出"])
    ws.append(["设备", 10000, 3000])
    result = extract_document(_bytes(wb), "预算.xlsx", "native_only")
    row = [u for u in result.units if u.kind == "table"][0]
    assert "经费 / 预算: 10000" in row.text and "经费 / 支出: 3000" in row.text
    assert "年度预算" in row.section


def test_xlsx_explicit_header_rows(monkeypatch):
    from openpyxl import Workbook
    monkeypatch.setattr(get_settings(), "rag_excel_header_rows", 2)
    wb = Workbook()
    ws = wb.active
    ws.append(["金额"])
    ws.append(["预算"])
    ws.append([20])
    result = extract(_bytes(wb), "a.xlsx")
    assert "金额 / 预算: 20" in result.text


def test_ppt_sections_survive_chunking():
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "仪器使用"
    slide.placeholders[1].text = "先登记设备。" * 200
    units = extract_document(_bytes(prs), "a.pptx", "native_only").units
    chunks = pack_units(units, source="a.pptx")
    assert len(chunks) > 1
    assert all("幻灯片 1" in c.text and "仪器使用" in c.text for c in chunks)
    assert all(not c.pages for c in chunks)  # Slide number must not masquerade as PDF page.


@pytest.mark.parametrize("extension", ["md", "markdown", "MD"])
def test_markdown_structures_and_nested_sections(extension):
    data = "# 设备\n\n## 外借\n\n- 先登记\n- 后审批\n\n```python\nprint('# 不是标题')\n```\n\n| 设备 | 管理员 |\n| --- | --- |\n| 示波器 | 张三 |\n\n## 归还\n\n检查配件。"
    result = extract_document(data.encode(), "a." + extension, "native_only")
    assert not result.error
    assert {u.kind for u in result.units} >= {"heading", "list", "code", "table"}
    assert next(u for u in result.units if u.kind == "code").section == ("设备", "外借")
    assert result.units[-1].section == ("设备", "归还")


def test_markdown_setext_heading_and_html_are_preserved():
    result = extract(b"Title\n=====\n\n<div>body</div>\n\n> quoted\n", "a.md")
    assert result.units[0].kind == "heading"
    assert any("<div>body</div>" in u.text for u in result.units)
    assert result.units[-1].section == ("Title",)


def test_exact_limit_has_no_redundant_tail():
    chunks = pack_units([TextUnit(1, "x" * 1000)], size=1000)
    assert len(chunks) == 1 and chunks[0].body == "x" * 1000


def test_long_text_full_coverage_without_contained_tail():
    text = "".join(chr(0x4E00 + i % 500) for i in range(1800))
    chunks = pack_units([TextUnit(3, text)], size=1000)
    rebuilt = chunks[0].body
    for before, after in zip(chunks, chunks[1:]):
        assert after.end > before.end
        rebuilt += after.body[max(0, before.end - after.start):]
    assert rebuilt == text
    assert all(c.pages == (3,) for c in chunks)
    assert all(token_bound(c.text) <= get_settings().rag_chunk_tokens for c in chunks)


def test_short_units_do_not_cross_sections():
    units = [TextUnit(1, "审批办法", section=("外借",)), TextUnit(2, "检查配件", section=("归还",))]
    chunks = pack_units(units, source="规章.pdf")
    assert len(chunks) == 2
    assert chunks[0].pages == (1,) and chunks[1].pages == (2,)
    assert all("规章.pdf" in c.text for c in chunks)


def test_multiple_explicit_chapters_inside_pdf_page_are_separated():
    chunks = pack_units([TextUnit(5, "第一章 外借\n先登记。\n第二章 归还\n检查配件。")])
    assert len(chunks) == 2
    assert chunks[0].section == ("第一章 外借",)
    assert chunks[1].section == ("第二章 归还",)
    assert all(c.pages == (5,) for c in chunks)


def test_plain_heading_detection_can_be_disabled(monkeypatch):
    monkeypatch.setattr(get_settings(), "rag_detect_plain_headings", False)
    chunks = pack_units([TextUnit(1, "第一章 外借\n正文。\n第二章 归还\n正文。")])
    assert len(chunks) == 1 and not chunks[0].section


def test_short_units_same_section_keep_pages_and_canonical_offsets():
    chunks = pack_units([TextUnit(1, "first", section=("A",)), TextUnit(2, "second", section=("A",))])
    assert len(chunks) == 1 and chunks[0].pages == (1, 2)
    assert chunks[0].end - chunks[0].start == len(chunks[0].body)


def test_sentence_boundary_preferred(monkeypatch):
    monkeypatch.setattr(get_settings(), "rag_chunk_tokens", 5000)
    text = "a" * 599 + "。" + "b" * 600
    chunks = pack_units([TextUnit(None, text)], size=1000, overlap=0)
    assert chunks[0].body.endswith("。")
    assert "".join(c.body for c in chunks) == text


def test_long_code_fragments_keep_fences_and_sections():
    data = "# 示例\n\n```python\n" + "print('hello')\n" * 300 + "```"
    chunks = pack_units(extract_document(data.encode(), "a.md", "native_only").units, source="a.md")
    code = [c for c in chunks if c.kind == "code"]
    assert len(code) > 1
    assert all(c.text.count("```") == 2 for c in code)
    assert all(len(c.text) <= get_settings().rag_code_chunk_chars for c in code)
    assert all("示例" in c.text for c in code)


def test_long_markdown_table_repeats_header():
    data = "# 台账\n\n| 编号 | 名称 |\n| --- | --- |\n" + "\n".join(f"| {i} | 设备{i} |" for i in range(200))
    chunks = pack_units(extract_document(data.encode(), "a.md", "native_only").units)
    rows = [c for c in chunks if c.kind == "table"]
    assert len(rows) > 1
    assert all("| 编号 | 名称 |" in c.text for c in rows)


def test_env_values_and_validation(monkeypatch):
    monkeypatch.setenv("RAG_CHUNK_CHARS", "500")
    monkeypatch.setenv("RAG_CHUNK_OVERLAP", "60")
    monkeypatch.setenv("RAG_EXCEL_HEADER_ROWS", "2")
    settings = Settings(_env_file=None)
    assert settings.rag_chunk_chars == 500 and settings.rag_chunk_overlap == 60
    assert settings.rag_excel_header_rows == 2
    with pytest.raises(ValueError):
        Settings(_env_file=None, rag_chunk_overlap=500)


def test_existing_file_parse_failure_does_not_publish(kb_env):
    kb_env.kb.files["a.pdf"] = make_pdf(pages=[{"text": "valid original"}])
    first = ingest.run_index_job(kb_env.kb, ingest.Scope(), ingest.Options())
    kb_env.kb.files["a.pdf"] = b"broken pdf"
    second = ingest.run_index_job(kb_env.kb, ingest.Scope(), ingest.Options())
    assert not second["published"] and second["failed_files"] == 1
    assert store.active_collection() == first["index_version"]
    assert store.collection_count(first["index_version"]) > 0


def test_markdown_end_to_end_index_retrieve_and_offsets(kb_env, monkeypatch):
    kb_env.kb.files["规章/a.md"] = "# 外借\n\n由设备管理员审批。".encode()
    summary = ingest.run_index_job(kb_env.kb, ingest.Scope(), ingest.Options())
    assert summary["published"]
    records, _ = kb_env.qdrant.scroll(summary["index_version"], with_payload=True)
    assert all(p.payload["chunking_version"].startswith("2-") for p in records)
    monkeypatch.setattr(retrieve, "get_embeddings", lambda: kb_env.embeddings)
    hits = retrieve.search_knowledge("设备审批", score_threshold=0)
    assert any("由设备管理员审批" in h["text"] for h in hits)
    assert all(h["ext"] == "md" for h in hits)


def _payload(index, body, start, section=None):
    return {"oss_key": "a.pdf", "source": "a.pdf", "ext": "pdf", "text": body, "body": body,
            "chunk_index": index, "content_sha256": "hash", "section_path": section or ["A"],
            "chunking_version": "2", "offset_version": "extracted-v2",
            "text_start": start, "text_end": start + len(body), "page_numbers": [index + 1]}


def test_neighbor_expansion_overlap_and_pages(kb_env, monkeypatch):
    collection = store.ensure_collection()
    a = _payload(0, "abcdefghijkl", 0)
    b = _payload(1, "ijklmnopqrst", 8)
    c = _payload(2, "不要混入另一章节", 20, ["B"])
    for i, p in enumerate((a, b, c)):
        store.upsert([store.make_point(str(__import__('uuid').uuid4()), kb_env.embeddings.embed_query(str(i)), p)])
    monkeypatch.setattr(retrieve, "get_embeddings", lambda: kb_env.embeddings)
    monkeypatch.setattr(store, "search", lambda *args, **kwargs: [SimpleNamespace(payload=b, score=0.9)])
    hits = retrieve.search_knowledge("q", top_k=1)
    assert len(hits) == 1 and "abcdefghijklmnopqrst" in hits[0]["text"]
    assert hits[0]["page_numbers"] == [1, 2]
    assert "不要混入" not in hits[0]["text"]


def test_neighbor_lookup_pins_collection_during_publish(kb_env, monkeypatch):
    initial = store.ensure_collection()
    p = _payload(1, "old evidence", 0)
    monkeypatch.setattr(retrieve, "get_embeddings", lambda: kb_env.embeddings)
    def search(*args, **kwargs):
        assert kwargs["collection"] == initial
        staged = store.create_staging()
        store.publish(staged, expected_count=0)
        return [SimpleNamespace(payload=p, score=0.9)]
    def neighbors(payload, *, collection, radius):
        assert collection == initial != store.active_collection()
        return []
    monkeypatch.setattr(store, "search", search)
    monkeypatch.setattr(store, "adjacent_chunks", neighbors)
    assert retrieve.search_knowledge("q")


def test_overfetch_dedup_and_context_budget(kb_env, monkeypatch):
    monkeypatch.setattr(get_settings(), "rag_context_chars", 256)
    monkeypatch.setattr(get_settings(), "rag_context_tokens", 256)
    monkeypatch.setattr(get_settings(), "rag_neighbor_radius", 0)
    a = _payload(0, "A" * 300, 0)
    duplicate = dict(a, chunk_index=1)
    b = dict(_payload(0, "B" * 300, 0), oss_key="b.pdf")
    monkeypatch.setattr(retrieve, "get_embeddings", lambda: kb_env.embeddings)
    def search(*args, **kwargs):
        assert kwargs["top_k"] > 2
        return [SimpleNamespace(payload=p, score=0.9) for p in (a, duplicate, b)]
    monkeypatch.setattr(store, "search", search)
    hits = retrieve.search_knowledge("q", top_k=2)
    assert sum(len(h["text"]) for h in hits) <= 256
    assert sum(token_bound(h["text"]) for h in hits) <= 256
    assert hits[0]["truncated"]
    assert "已截断" in retrieve.format_hits(hits)


def test_same_text_different_files_remain_distinct():
    a = _payload(0, "same", 0)
    assert not retrieve._redundant(a, dict(a, oss_key="b.pdf"))


def test_legacy_index_payload_remains_readable(kb_env, monkeypatch):
    monkeypatch.setattr(retrieve, "get_embeddings", lambda: kb_env.embeddings)
    payload = {"oss_key": "legacy.txt", "text": "旧版正文", "source": "legacy.txt"}
    monkeypatch.setattr(store, "search", lambda *a, **kw: [SimpleNamespace(payload=payload, score=0.9)])
    result = retrieve.search_knowledge("q")
    assert result[0]["text"] == "旧版正文" and result[0]["page_numbers"] == []


def test_neighbor_failure_keeps_primary(kb_env, monkeypatch):
    monkeypatch.setattr(retrieve, "get_embeddings", lambda: kb_env.embeddings)
    monkeypatch.setattr(store, "search", lambda *a, **kw: [SimpleNamespace(payload=_payload(0, "primary", 0), score=0.9)])
    def fail(*args, **kwargs):
        raise RuntimeError("neighbor query unavailable")
    monkeypatch.setattr(store, "adjacent_chunks", fail)
    assert "primary" in retrieve.search_knowledge("q")[0]["text"]


def test_transitive_adjacent_anchors_merge_in_source_order(kb_env, monkeypatch):
    monkeypatch.setattr(get_settings(), "rag_neighbor_radius", 0)
    monkeypatch.setattr(retrieve, "get_embeddings", lambda: kb_env.embeddings)
    a = _payload(0, "abcdefghij", 0)
    b = _payload(1, "ijklmnopqr", 8)
    c = _payload(2, "qrstuvwxyz", 16)
    monkeypatch.setattr(store, "search", lambda *args, **kwargs: [
        SimpleNamespace(payload=p, score=score) for p, score in ((a, .9), (c, .8), (b, .7))])
    hits = retrieve.search_knowledge("q", top_k=3)
    assert len(hits) == 1 and "abcdefghijklmnopqrstuvwxyz" in hits[0]["text"]
    assert hits[0]["chunk_indices"] == [0, 1, 2] and hits[0]["score"] == .9


def test_wide_table_does_not_drop_original_content():
    header = "| " + " | ".join("列名" + str(i) for i in range(200)) + " |"
    separator = "| " + " | ".join("---" for _ in range(200)) + " |"
    raw = header + "\n" + separator + "\n| 数据 |"
    chunks = pack_units([TextUnit(None, raw, section=("台账",), kind="table")])
    assert len(chunks) > 1
    assert "".join(c.body for c in chunks) == raw
    assert all(len(c.text) <= get_settings().rag_table_chunk_chars for c in chunks)
    assert all(token_bound(c.text) <= get_settings().rag_chunk_tokens for c in chunks)


def test_utf8_bom_does_not_hide_markdown_heading():
    result = extract("# 标题\n\n正文".encode("utf-8-sig"), "a.md")
    assert result.units[0].kind == "heading" and result.units[-1].section == ("标题",)


def test_formatted_context_budget_includes_metadata(monkeypatch):
    monkeypatch.setattr(get_settings(), "rag_context_chars", 256)
    monkeypatch.setattr(get_settings(), "rag_context_tokens", 256)
    hits = [{"source": "source.pdf", "section_path": ["长标题" * 1000], "text": "正文", "page_numbers": [1]}]
    text = retrieve.format_hits(hits)
    assert len(text) <= 256 and token_bound(text) <= 256
    assert "已截断" in text and hits[0]["truncated"]


def test_citations_union_pages_and_sections_without_mutating_artifact():
    from langchain_core.messages import ToolMessage
    from app.graph.trace import used_sources
    from app.api.routes.chat import _sign_sources
    a = dict(_payload(0, "first", 0), score=0.9, chunk_indices=[0])
    b = dict(_payload(1, "second", 6, ["B"]), score=0.8, chunk_indices=[1])
    message = ToolMessage(content="hits", name="search_knowledge", tool_call_id="1", artifact=[a, b])
    result = used_sources([message])
    assert result[0]["page_numbers"] == [1, 2] and result[0]["section_paths"] == [["A"], ["B"]]
    assert a["page_numbers"] == [1]
    citation = _sign_sources(result)[0]
    assert citation.page_numbers == [1, 2] and citation.chunk_indices == [0, 1]
