"""RAG 抽取器:按文件扩展名把原始字节抽成纯文本(摄取 2e 的原料层)。

纯函数,吃 `(data: bytes, filename: str)` 吐 `ExtractResult`,不连 OSS / Qdrant。
设计原则:
- 坏文件 / 不支持的类型**不抛异常**,把原因写进 `error` 字段,让整批摄取不中断
  (同「零命中不崩」原则)。
- 各解析库**惰性 import**(未装也能导入本模块,真用到才报错)。
- 扫描件 PDF(无文本层)与图片(png/jpg)只标 `needs_ocr=True`,真 OCR 在 2d 接。
- `.doc` 老格式经 LibreOffice headless 转成 .docx 字节后再抽取(见 doc_convert;未装则标
  unsupported);`.xls` 老格式需 xlrd(未装则标 unsupported);`.zip` v1 不解包。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from app.rag import doc_convert
from app.rag.structure import ExtractUnit


@dataclass
class ExtractResult:
    text: str = ""                       # 抽出的正文(可能为空)
    blocks: list[str] = field(default_factory=list)  # 自然块(页/段/表行),供 2e 切块;无则 [text]
    needs_ocr: bool = False              # 图片 / 扫描 PDF → True,交 2d OCR
    ext: str = ""                        # 小写扩展名(无点)
    extractor: str = ""                  # 实际用的抽取器:pypdf / python-docx / openpyxl / ...
    error: str | None = None             # unsupported / skipped / 解析失败(不抛,标记)
    units: list[ExtractUnit] = field(default_factory=list)


# ---- 小工具 ----

def _ext(filename: str) -> str:
    """取小写扩展名(无点);无扩展名返回空串。"""
    name = (filename or "").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    if "." not in name:
        return ""
    return name.rsplit(".", 1)[-1].lower()


def _decode(data: bytes) -> str:
    """文本字节解码:utf-8-sig → gbk → gb18030，最后兜底 replace。"""
    for enc in ("utf-8-sig", "gbk", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell(v) -> str:
    """表格单元格值归一成字符串。"""
    if v is None:
        return ""
    return str(v).strip()


def _xls_num(v):
    """BIFF(.xls)里的数字一律按浮点存:整数值归一成 int,免得文本里出现「1200.0」。

    只影响显示成什么样;小数原样保留(1200.5 不动),openpyxl 那条路径本来就是 int。
    """
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


# ---- 各类型抽取器(签名统一:(data, ext) -> ExtractResult)----

def _from_txt(data: bytes, ext: str) -> ExtractResult:
    text = _decode(data).strip()
    blocks = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return ExtractResult(text=text, blocks=blocks, ext=ext, extractor="txt")


def _from_markdown(data: bytes, ext: str) -> ExtractResult:
    """用 CommonMark 源行范围保留列表、围栏代码和表格，标题路径单独传播。"""
    from markdown_it import MarkdownIt

    text = _decode(data).strip()
    lines = text.splitlines()
    tokens = MarkdownIt("commonmark").enable("table").parse(text)
    units: list[ExtractUnit] = []
    headings: list[tuple[int, str]] = []
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open":
            level = int(token.tag[1:])
            title = tokens[i + 1].content
            headings = [h for h in headings if h[0] < level] + [(level, title)]
            units.append(ExtractUnit(title, tuple(h[1] for h in headings), "heading"))
            i += 3
            continue
        if token.level == 0 and token.map and token.nesting >= 0:
            start, end = token.map
            body = "\n".join(lines[start:end]).strip()
            kind = {"fence": "code", "code_block": "code", "table_open": "table",
                    "bullet_list_open": "list", "ordered_list_open": "list"}.get(token.type, "paragraph")
            if body:
                units.append(ExtractUnit(body, tuple(h[1] for h in headings), kind))
            if token.nesting == 1:
                closing = token.type.replace("_open", "_close")
                i += 1
                while i < len(tokens) and not (tokens[i].type == closing and tokens[i].level == 0):
                    i += 1
        i += 1
    return ExtractResult(text=text, blocks=[u.text for u in units], units=units,
                         ext=ext, extractor="markdown-it")


def _from_pdf(data: bytes, ext: str) -> ExtractResult:
    """电子版 PDF 抽文本层;无文本层(扫描版)→ needs_ocr,交 2d。"""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages: list[str] = []
        for page in reader.pages:
            t = (page.extract_text() or "").strip()
            if t:
                pages.append(t)
        text = "\n\n".join(pages).strip()
        if not text:
            return ExtractResult(needs_ocr=True, ext=ext, extractor="pypdf")
        return ExtractResult(text=text, blocks=pages, ext=ext, extractor="pypdf")
    except Exception as exc:
        return ExtractResult(ext=ext, extractor="pypdf", error=f"pdf 解析失败: {exc}")


def _from_docx(data: bytes, ext: str) -> ExtractResult:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        from docx.oxml.ns import qn
        import re

        units: list[ExtractUnit] = []
        headings: list[tuple[int, str]] = []
        table_number = 0
        # 按 body XML 顺序遍历，保留段落与表格的交错顺序。
        for child in doc.element.body:
            section = tuple(h[1] for h in headings)
            if child.tag == qn("w:p"):
                para = Paragraph(child, doc)
                t = para.text.strip()
                if not t:
                    continue
                match = re.search(r"(?:Heading|标题)\s*(\d+)", para.style.name, re.I)
                if match:
                    level = int(match.group(1))
                    headings = [h for h in headings if h[0] < level] + [(level, t)]
                    section = tuple(h[1] for h in headings)
                units.append(ExtractUnit(t, section, "heading" if match else "paragraph"))
            elif child.tag == qn("w:tbl"):
                table_number += 1
                table = Table(child, doc)
                rows = [[c.text.strip() for c in row.cells] for row in table.rows]
                header = rows[0] if rows else []
                for row in rows:
                    # 保留位置和空单元格；后续行携带列名，便于独立召回。
                    line = " | ".join(row)
                    if any(row):
                        labels = "; ".join(f"{h}: {v}" for h, v in zip(header, row) if h and v)
                        units.append(ExtractUnit(line + ("\n" + labels if row != header and labels else ""),
                                                 section + (f"表格 {table_number}", "表头: " + " | ".join(header)), "table"))
        blocks = [u.text for u in units]
        text = "\n".join(blocks).strip()
        return ExtractResult(text=text, blocks=blocks, units=units, ext=ext, extractor="python-docx")
    except Exception as exc:
        return ExtractResult(ext=ext, extractor="python-docx", error=f"docx 解析失败: {exc}")


def _from_xlsx(data: bytes, ext: str) -> ExtractResult:
    """xlsx 按行转带表头语义文本(每行 `列名: 值`),一 sheet 一段。"""
    try:
        from openpyxl import load_workbook

        from app.config import get_settings
        wb = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
        blocks: list[str] = []
        units: list[ExtractUnit] = []
        for ws in wb.worksheets:
            rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
            if not rows or not any(any(row) for row in rows):
                continue
            merges = list(ws.merged_cells.ranges)
            # 合并区域展开只用于表头解析，数据区域不凭空复制值。
            expanded = [list(row) for row in rows[:4]]
            for area in merges:
                if area.min_row <= 4:
                    value = _cell(ws.cell(area.min_row, area.min_col).value)
                    for r in range(area.min_row - 1, min(area.max_row, len(expanded))):
                        for c in range(area.min_col - 1, area.max_col):
                            expanded[r][c] = value
            configured = get_settings().rag_excel_header_rows
            title_rows = 0
            if not configured and any(a.min_row == 1 and a.max_row == 1 and a.min_col == 1
                                      and a.max_col == ws.max_column for a in merges) and len(rows) > 1:
                title_rows = 1
            count = configured or max([title_rows + 1] + [min(a.max_row + (a.max_col > a.min_col), 4)
                        for a in merges if a.min_row == title_rows + 1])
            count = min(max(1, count), 4, len(rows))
            header = []
            for c in range(ws.max_column):
                parts = list(dict.fromkeys(row[c] for row in expanded[title_rows:count] if row[c]))
                header.append(" / ".join(parts))
            section = (f"工作表: {ws.title}",)
            if title_rows:
                section += (rows[0][0],)
            blocks.append(f"[表: {ws.title}]")
            units.append(ExtractUnit(blocks[-1], section, "heading"))
            for row in rows[count:]:
                if row is None:
                    continue
                pairs: list[str] = []
                for h, v in zip(header, row):
                    cv = _cell(v)
                    if cv == "":
                        continue
                    pairs.append(f"{h}: {cv}" if h else cv)
                if pairs:
                    blocks.append("; ".join(pairs))
                    units.append(ExtractUnit(blocks[-1], section + ("表头: " + " | ".join(header),), "table"))
        wb.close()
        text = "\n".join(blocks).strip()
        return ExtractResult(text=text, blocks=blocks, units=units, ext=ext, extractor="openpyxl")
    except Exception as exc:
        return ExtractResult(ext=ext, extractor="openpyxl", error=f"xlsx 解析失败: {exc}")


def _from_xls(data: bytes, ext: str) -> ExtractResult:
    """.xls 老格式:走 xlrd;环境里没装(未 pip install -r requirements.txt)则标 unsupported。"""
    try:
        import xlrd
    except Exception:
        return ExtractResult(ext=ext, extractor="none",
                             error="unsupported: .xls 老格式需 xlrd 解析,当前环境未安装(安装 requirements.txt 里的 xlrd>=2.0 后重试)")
    try:
        from app.config import get_settings
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
        blocks: list[str] = []
        units: list[ExtractUnit] = []
        for sheet in book.sheets():
            if sheet.nrows == 0:
                continue
            merged = getattr(sheet, "merged_cells", [])
            expanded = [[_cell(_xls_num(sheet.cell_value(r, c))) for c in range(sheet.ncols)]
                        for r in range(min(sheet.nrows, 4))]
            for r0, r1, c0, c1 in merged:
                if r0 < len(expanded):
                    value = expanded[r0][c0]
                    for r in range(r0, min(r1, len(expanded))):
                        for c in range(c0, c1):
                            expanded[r][c] = value
            configured = get_settings().rag_excel_header_rows
            title_rows = int(not configured and sheet.nrows > 1
                             and (0, 1, 0, sheet.ncols) in merged)
            count = configured or max([title_rows + 1] + [min(r1 + (c1 - c0 > 1), 4)
                                      for r0, r1, c0, c1 in merged if r0 == title_rows])
            count = min(max(1, count), 4, sheet.nrows)
            header = [" / ".join(dict.fromkeys(row[c] for row in expanded[title_rows:count] if row[c]))
                      for c in range(sheet.ncols)]
            blocks.append(f"[表: {sheet.name}]")
            section = (f"工作表: {sheet.name}",)
            if title_rows:
                section += (expanded[0][0],)
            units.append(ExtractUnit(blocks[-1], section, "heading"))
            for r in range(count, sheet.nrows):
                pairs: list[str] = []
                for c in range(sheet.ncols):
                    cv = _cell(_xls_num(sheet.cell_value(r, c)))
                    if cv == "":
                        continue
                    pairs.append(f"{header[c]}: {cv}" if header[c] else cv)
                if pairs:
                    blocks.append("; ".join(pairs))
                    units.append(ExtractUnit(blocks[-1], section + ("表头: " + " | ".join(header),), "table"))
        text = "\n".join(blocks).strip()
        return ExtractResult(text=text, blocks=blocks, units=units, ext=ext, extractor="xlrd")
    except Exception as exc:
        return ExtractResult(ext=ext, extractor="xlrd", error=f"xls 解析失败: {exc}")


def _from_pptx(data: bytes, ext: str) -> ExtractResult:
    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        blocks: list[str] = []
        units: list[ExtractUnit] = []
        for i, slide in enumerate(prs.slides, 1):
            texts: list[str] = []
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
            section = (f"幻灯片 {i}",) + ((title,) if title else ())
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs).strip()
                        if line:
                            texts.append(line)
                            units.append(ExtractUnit(line, section))
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        rl = " | ".join(c for c in cells if c)
                        if rl:
                            texts.append(rl)
                            units.append(ExtractUnit(rl, section, "table"))
            if texts:
                blocks.append(f"[幻灯片 {i}]")
                blocks.extend(texts)
        text = "\n".join(blocks).strip()
        return ExtractResult(text=text, blocks=blocks, units=units, ext=ext, extractor="python-pptx")
    except Exception as exc:
        return ExtractResult(ext=ext, extractor="python-pptx", error=f"pptx 解析失败: {exc}")


def _from_image(data: bytes, ext: str) -> ExtractResult:
    """图片:全部交 2d OCR;本步只占位标记。"""
    return ExtractResult(needs_ocr=True, ext=ext, extractor="ocr-pending")


def _from_doc(data: bytes, ext: str) -> ExtractResult:
    """.doc 老格式:先经 LibreOffice 转成 .docx 字节,再走同一条 docx 抽取路径。

    转换失败(未装 / 超时 / 坏文件)按 unsupported / 转换失败记进 error,不抛。
    """
    docx, err = doc_convert.doc_to_docx(data)
    if docx is None:
        return ExtractResult(ext=ext, extractor="none", error=err or "doc 转换失败")
    res = _from_docx(docx, ext)
    res.extractor = "libreoffice+python-docx"
    if res.error:
        res.error = f"doc 转换成功但解析失败: {res.error}"
    return res


def _skip_zip(data: bytes, ext: str) -> ExtractResult:
    return ExtractResult(ext=ext, extractor="none", error="skipped: .zip 压缩包 v1 不解包")


_DISPATCH = {
    "pdf": _from_pdf,
    "docx": _from_docx,
    "xlsx": _from_xlsx,
    "xls": _from_xls,
    "pptx": _from_pptx,
    "txt": _from_txt,
    "md": _from_markdown,
    "markdown": _from_markdown,
    "png": _from_image,
    "jpg": _from_image,
    "jpeg": _from_image,
    "doc": _from_doc,
    "zip": _skip_zip,
}

# 支持(能抽出文本或明确交 OCR)的扩展名,供上层预判 / 提示。
SUPPORTED_EXTS = frozenset(_DISPATCH)


def extract(data: bytes, filename: str) -> ExtractResult:
    """按扩展名分派抽取;未知扩展名标 unsupported(不抛)。"""
    ext = _ext(filename)
    fn = _DISPATCH.get(ext)
    if fn is None:
        return ExtractResult(ext=ext, extractor="none", error=f"unsupported: 未知扩展名 .{ext or '(无)'}")
    return fn(data, ext)
